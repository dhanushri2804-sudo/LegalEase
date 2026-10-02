from __future__ import annotations

import io
from typing import Any

from docx import Document
from docx.shared import Pt
from fpdf import FPDF
from fpdf.enums import XPos, YPos


def _normalize_document(document: Any) -> dict[str, Any]:
    if isinstance(document, dict):
        return document

    return {
        "title": getattr(document, "title", "Legal Document"),
        "document_type": getattr(document, "document_type", "General Agreement"),
        "party_details": getattr(document, "party_details", ""),
        "terms": getattr(document, "terms", ""),
        "effective_date": getattr(document, "effective_date", ""),
        "jurisdiction": getattr(document, "jurisdiction", ""),
        "generated_content": getattr(document, "generated_content", ""),
    }


def format_txt(document: Any) -> str:
    record = _normalize_document(document)
    title = record.get("title") or "Legal Document"
    doc_type = record.get("document_type") or "General Agreement"
    content = record.get("generated_content") or ""

    output = [
        title,
        "=" * len(title),
        f"Document Type: {doc_type}",
        "",
        content.strip(),
    ]
    return "\n".join(output) + "\n"


def format_docx(document: Any, company_name: str | None = None, logo_path: str | None = None) -> bytes:
    record = _normalize_document(document)
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    title_para = doc.add_paragraph()
    title_run = title_para.add_run(record.get("title") or "Legal Document")
    title_run.bold = True
    title_run.font.name = "Times New Roman"
    title_run.font.size = Pt(22)

    if company_name:
        doc.add_paragraph(f"Company: {company_name}")

    doc.add_paragraph(f"Document Type: {record.get('document_type') or 'General Agreement'}")
    doc.add_paragraph(f"Effective Date: {record.get('effective_date') or '[Effective Date]'}")
    if record.get("jurisdiction"):
        doc.add_paragraph(f"Jurisdiction: {record.get('jurisdiction')}")

    content = record.get("generated_content") or ""
    sections = content.splitlines()
    for line in sections:
        paragraph = doc.add_paragraph()
        if line.strip():
            paragraph.alignment = 3
            run = paragraph.add_run(line)
            run.font.name = "Times New Roman"
            run.font.size = Pt(12)

    if logo_path:
        try:
            doc.add_paragraph()
            doc.add_picture(logo_path, width=120)
        except Exception as exc:
            raise ValueError("The provided company logo could not be added to the document.") from exc

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def format_pdf(document: Any, company_name: str | None = None) -> bytes:
    record = _normalize_document(document)
    content = (record.get("generated_content") or "").splitlines()
    title = record.get("title") or "Legal Document"

    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.set_font("Times", "B", 18)
    pdf.cell(0, 12, title, new_x=XPos.LMARGIN, new_y=YPos.NEXT)

    if company_name:
        pdf.set_font("Times", "", 11)
        pdf.cell(0, 8, f"Company: {company_name}", new_x=XPos.LMARGIN, new_y=YPos.NEXT)

    pdf.set_font("Times", "", 11)
    pdf.ln(4)

    for line in content:
        if not line.strip():
            pdf.ln(4)
            continue
        if len(line) > 95:
            chunks = [line[i:i + 95] for i in range(0, len(line), 95)]
            for chunk in chunks:
                pdf.multi_cell(0, 7, chunk)
        else:
            pdf.multi_cell(0, 7, line)

    pdf.set_y(-15)
    pdf.set_font("Times", "I", 8)
    pdf.cell(0, 10, f"Page {pdf.page_no()}", align="C")

    return bytes(pdf.output())
